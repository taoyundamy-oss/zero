"""C3.2 验证种子：生成 docx/xlsx/png 测试文件并登记 file_id（在 worker 容器内运行）。

用法：
  docker cp /opt/zero-cloud/scripts/seed_c32.py zero-worker:/tmp/seed_c32.py
  docker exec zero-worker python /tmp/seed_c32.py
输出：每文件一行 `<name> <登记JSON>`（file_id 用于 MCP tools/call）。
"""
import json
import sys
from pathlib import Path

sys.path.insert(0, "/app")
from tasks.file import _register_file  # noqa: E402（复用登记：sha256/TTL/Z6 同构）

out = Path("/data/zero/test")
out.mkdir(parents=True, exist_ok=True)

from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from PIL import Image  # noqa: E402

d = Document()
d.add_paragraph("Zero C3.2 docx_read 测试段落一：项目代号 Fable。")
d.add_paragraph("测试段落二：云服文件处理管线全通。")
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text = "键", "值"
t.cell(1, 0).text, t.cell(1, 1).text = "轨", "C3.2"
p_docx = out / "c32-test.docx"
d.save(p_docx)

wb = Workbook()
ws = wb.active
ws.title = "数据"
ws.append(["名称", "数量", "备注"])
ws.append(["苹果", 3, "fresh"])
ws.append(["香蕉", 5, "ripe"])
p_xlsx = out / "c32-test.xlsx"
wb.save(p_xlsx)

im = Image.new("RGB", (800, 600), (30, 144, 255))
p_png = out / "c32-test.png"
im.save(p_png)

MIME = {
    p_docx: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    p_xlsx: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    p_png: "image/png",
}
for path, mime in MIME.items():
    print(path.name, json.dumps(_register_file(path, mime), ensure_ascii=False))
